"""Resume text extraction (PDF / DOCX / TXT) and format-tolerant structured parsing."""
from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass, field, fields
from datetime import date

from app.errors import ResumeParseError
from resume.lexicon import find_terms

PARSER_VERSION = 1

# ----------------------------------------------------------------------------------------------- text extraction
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", "replace")


def _fix_pdf_text(t: str) -> str:
    t = re.sub(r"(\w)-\n(\w)", r"\1\2", t)
    lines = [ln.rstrip() for ln in t.replace("\r", "\n").split("\n")]
    return "\n".join(ln for ln in lines if not re.fullmatch(r"\s*(?:page\s*)?\d{1,3}(?:\s*(?:/|of)\s*\d{1,3})?\s*", ln, re.I))


def extract_text(data: bytes, filename: str, cancel=None, progress=None) -> str:
    """`cancel`: object with is_set(); `progress(done, total)` is called per PDF page."""
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if not data:
        raise ResumeParseError(f"{filename}: the file is empty.")
    if ext in (".txt", ".md", ".text"):
        text = _decode(data)
    elif ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ResumeParseError(f"{filename}: the PDF is password-protected.")
            pages, total = [], len(reader.pages)
            for i, page in enumerate(reader.pages):
                if cancel is not None and cancel.is_set():
                    from app.errors import Cancelled
                    raise Cancelled()
                pages.append(page.extract_text() or "")
                if progress:
                    progress(i + 1, total)
            text = _fix_pdf_text("\n".join(pages))
        except ResumeParseError:
            raise
        except Exception as e:
            if getattr(e, "cancelled", False):
                raise
            raise ResumeParseError(f"{filename}: not a readable PDF ({type(e).__name__}).") from e
    elif ext == ".docx":
        try:
            import docx
            d = docx.Document(io.BytesIO(data))
            lines = []
            for p in d.paragraphs:
                t = p.text.strip()
                if t:
                    lines.append(("- " + t) if "list" in (p.style.name or "").lower() else t)
            for tbl in d.tables:
                for row in tbl.rows:
                    cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
                    if cells:
                        lines.append(" | ".join(cells))
            text = "\n".join(lines)
        except Exception as e:
            raise ResumeParseError(f"{filename}: not a readable DOCX ({type(e).__name__}).") from e
    else:
        raise ResumeParseError(f"{filename}: unsupported file type '{ext or 'unknown'}'. Use PDF, DOCX or TXT.")
    if len(text.split()) < 5:
        hint = " It may be a scanned image (OCR is not supported)." if ext == ".pdf" else ""
        raise ResumeParseError(f"{filename}: no readable text found.{hint}")
    return text


# ----------------------------------------------------------------------------------------------- data model
@dataclass
class ParsedResume:
    name: str = ""
    summary: str = ""
    skills: list = field(default_factory=list)
    programming_languages: list = field(default_factory=list)
    frameworks: list = field(default_factory=list)
    tools: list = field(default_factory=list)
    technologies: list = field(default_factory=list)
    education: list = field(default_factory=list)
    experience: list = field(default_factory=list)
    companies: list = field(default_factory=list)
    job_titles: list = field(default_factory=list)
    projects: list = field(default_factory=list)
    responsibilities: list = field(default_factory=list)
    achievements: list = field(default_factory=list)
    certifications: list = field(default_factory=list)
    years_of_experience: float = 0.0
    sections_found: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict | None) -> "ParsedResume":
        d = d or {}
        names = {f.name: f for f in fields(cls)}
        kw = {}
        for k, v in d.items():
            if k in names:
                kw[k] = v
        return cls(**kw)


# ----------------------------------------------------------------------------------------------- parsing helpers
_SECTIONS = {
    "summary": r"(?:professional |career |executive )?(?:summary|profile|objective|about(?: me)?|overview|career objective)",
    "skills": r"(?:(?:technical|key|core|it|professional) )?(?:skills?|competenc(?:y|ies)|technolog(?:y|ies)|tech stack|expertise|tools(?: (?:&|and) technologies)?|skills (?:&|and) tools)(?: summary)?",
    "experience": r"(?:(?:professional|work|relevant|industry|employment) )?(?:experience|history)|employment(?: history)?|work history|career history|internships?",
    "education": r"education(?:al)?(?: (?:background|qualifications?))?|academics?|academic (?:background|qualifications?)|qualifications?",
    "projects": r"(?:(?:personal|academic|key|selected|notable|major|side|technical) )?projects?",
    "certifications": r"certifications?(?: (?:&|and) (?:licenses?|courses?|training))?|licenses?(?: (?:&|and) certifications?)?|courses?|training|certificates?",
    "achievements": r"achievements?|awards?(?: (?:&|and) honou?rs?)?|honou?rs?|accomplishments?|awards? (?:&|and) achievements?",
}
_SEC_RE = {k: re.compile(r"^\W*(?:" + v + r")\s*[:\-–—]?\W*$", re.I) for k, v in _SECTIONS.items()}
_INLINE_SEC = re.compile(r"^\W*(" + "|".join(f"(?P<{k}>{v})" for k, v in _SECTIONS.items()) + r")\s*[:\-–—]\s*(?P<rest>\S.*)$", re.I)
_BULLET = re.compile(r"^\s*(?:[-*•●▪◦■➢►–·]|\d{1,2}[.)])\s+")
_MONTH = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
_DATE = rf"(?:{_MONTH}\.?,?\s+\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
_RANGE = re.compile(rf"(?P<s>{_DATE})\s*(?:-|–|—|to|until)\s*(?P<e>{_DATE}|present|current|now|ongoing|till date|today)", re.I)
_SINCE = re.compile(rf"\b(?:since|from)\s+(?P<s>{_DATE})\b", re.I)
_CONTACT = re.compile(r"(@|https?://|www\.|linkedin|github\.com|\+?\d[\d\s().-]{7,}\d)", re.I)
_TITLE_WORDS = re.compile(r"\b(engineer|developer|manager|analyst|intern|internship|lead|consultant|architect|scientist|director|associate|"
                          r"specialist|administrator|designer|officer|head|vp|president|founder|co-founder|programmer|researcher|"
                          r"trainee|executive|coordinator|technician|assistant|professor|lecturer|teacher|tester|qa|sde|swe|devops|"
                          r"freelancer|contractor|supervisor|principal|staff|senior|junior|fellow)\b", re.I)
_ACHIEVE = re.compile(r"(\d|%|\b(?:improv|reduc|increas|achiev|led|lead|saved|won|award|boost|decreas|cut|grew)\w*)", re.I)


def _norm_line(ln: str) -> str:
    return re.sub(r"[ \t]+", " ", ln.replace(" ", " ")).strip()


def _section_of(line: str) -> str | None:
    s = _norm_line(line)
    if not s or len(s) > 60 or _BULLET.match(s):
        return None
    core = re.sub(r"^[\W_]+|[\W_]+$", "", s)
    for k, rx in _SEC_RE.items():
        if rx.match(core + ":") or rx.match(core):
            if core.isupper() or s.endswith(":") or re.match(r"^[=\-_*#~\s]+|[=\-_*#~\s]+$", s) or core.istitle() or len(core.split()) <= 4:
                return k
    return None


def _to_date(s: str, end: bool = False) -> date | None:
    s = s.strip().lower().rstrip(".,")
    if s in ("present", "current", "now", "ongoing", "till date", "today"):
        return date.today()
    m = re.match(rf"({_MONTH})\.?,?\s+(\d{{4}})", s)
    mon = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
    if m:
        return date(int(m.group(2)), mon[m.group(1)[:3]], 1)
    m = re.match(r"(\d{1,2})/(\d{4})", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return date(int(m.group(2)), int(m.group(1)), 1)
    m = re.match(r"(\d{4})$", s)
    if m:
        return date(int(m.group(1)), 12 if end else 1, 1)
    return None


def _months(a: date, b: date) -> int:
    return max(0, (b.year - a.year) * 12 + b.month - a.month + (1 if b >= a else 0))


def _split_sections(lines: list[str]) -> tuple[list[str], dict[str, list[str]], list[str]]:
    head: list[str] = []
    secs: dict[str, list[str]] = {}
    order: list[str] = []
    cur: str | None = None
    for raw in lines:
        ln = raw.rstrip()
        if re.fullmatch(r"\s*[=\-_*~#]{3,}\s*", ln):
            continue
        sec = _section_of(ln)
        if sec:
            cur = sec
            secs.setdefault(sec, []); order.append(sec) if sec not in order else None
            continue
        m = _INLINE_SEC.match(_norm_line(ln))
        if m and not _BULLET.match(ln):
            key = next(k for k in _SECTIONS if m.group(k))
            if key in ("skills", "summary", "education", "certifications", "achievements", "projects"):
                cur = key
                secs.setdefault(key, []).append(m.group("rest")); order.append(key) if key not in order else None
                continue
        (secs[cur] if cur else head).append(ln)
    return head, secs, order


def _split_list(s: str) -> list[str]:
    parts = re.split(r"[,;|•●▪·]|\s{2,}|\s/\s", s)
    return [p.strip(" .-–") for p in parts if 1 < len(p.strip(" .-–")) <= 50]


def _guess_name(head: list[str]) -> str:
    for ln in head[:8]:
        s = _norm_line(ln)
        if not s or _CONTACT.search(s) or any(ch.isdigit() for ch in s):
            continue
        words = s.split()
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z][A-Za-z.'\-]*", w) for w in words) \
                and not _TITLE_WORDS.search(s) and not _section_of(s):
            return " ".join(w.capitalize() if w.isupper() else w for w in words)
    return ""


def _split_title_company(text: str) -> tuple[str, str]:
    text = _RANGE.sub("", text)
    text = re.sub(r"\(\s*\)|\s{2,}", " ", text).strip(" |,-–—@·")
    parts = [p.strip() for p in re.split(r"\s+at\s+|\s*@\s*|\s*\|\s*|\s+[-–—]\s+|,\s*", text) if p.strip()]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return (parts[0], "") if _TITLE_WORDS.search(parts[0]) else ("", parts[0])
    a, b = parts[0], parts[1]
    if _TITLE_WORDS.search(a) and not _TITLE_WORDS.search(b):
        return a, b
    if _TITLE_WORDS.search(b) and not _TITLE_WORDS.search(a):
        return b, a
    return a, b


def _parse_experience(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    cur: dict | None = None
    pending_header: list[str] = []
    for raw in lines:
        s = _norm_line(raw)
        if not s:
            continue
        rng = _RANGE.search(s) or None
        is_b = bool(_BULLET.match(raw))
        if rng and not is_b and len(s) <= 140:
            hdr = _RANGE.sub("", s).strip(" |,-–—()")
            title, company = _split_title_company(hdr)
            if not (title or company) and pending_header:
                title, company = _split_title_company(" | ".join(pending_header))
            elif not company and pending_header and not _BULLET.match(pending_header[-1]):
                prev_t, prev_c = _split_title_company(pending_header[-1])
                title, company = title or prev_t, company or prev_c or pending_header[-1]
            cur = {"title": title, "company": company, "start": rng.group("s"), "end": rng.group("e"), "months": 0,
                   "location": "", "bullets": [], "responsibilities": [], "achievements": [], "technologies": []}
            a, b = _to_date(rng.group("s")), _to_date(rng.group("e"), True)
            if a and b:
                cur["months"] = _months(a, b)
            entries.append(cur)
            pending_header = []
            continue
        if is_b:
            b = _BULLET.sub("", s).strip()
            if cur is not None:
                cur["bullets"].append(b)
            pending_header = []
            continue
        if cur is not None and not _TITLE_WORDS.search(s) and (len(s) > 60 or s.endswith(".") or cur["bullets"] or cur is entries[-1] and entries and not pending_header and len(s.split()) > 6):
            if cur["bullets"] and len(s) < 90 and _TITLE_WORDS.search(s):
                pending_header.append(s)
            else:
                cur["bullets"].append(s) if len(s.split()) > 3 else pending_header.append(s)
            continue
        pending_header = (pending_header + [s])[-2:]
        if cur is not None and (len(s) > 60):
            cur["bullets"].append(s)
    for e in entries:
        for b in e["bullets"]:
            e["responsibilities"].append(b)
            if _ACHIEVE.search(b):
                e["achievements"].append(b)
        terms = find_terms(" ".join(e["bullets"]))
        e["technologies"] = list(dict.fromkeys(terms["programming_languages"] + terms["frameworks"] + terms["tools"]))
    return [e for e in entries if e["title"] or e["company"]]


def _parse_projects(lines: list[str]) -> list[dict]:
    projs: list[dict] = []
    cur: dict | None = None
    for raw in lines:
        s = _norm_line(raw)
        if not s:
            continue
        if _BULLET.match(raw):
            b = _BULLET.sub("", s).strip()
            if cur is None:
                cur = {"name": b[:60], "description": "", "technologies": [], "bullets": []}
                projs.append(cur)
            else:
                cur["bullets"].append(b)
            continue
        is_head = len(s) <= 90 and not s.endswith(".") and (cur is None or cur["bullets"] or cur["description"] or len(s.split()) <= 8)
        tech = re.match(r"^(?:tech(?:nolog(?:y|ies))?(?: stack)?|tools?|built with|stack)\s*[:\-]\s*(.+)$", s, re.I)
        if tech and cur is not None:
            cur["technologies"] += _split_list(tech.group(1))
        elif is_head and (cur is None or cur["bullets"] or cur["description"]):
            m = re.split(r"\s+[–—|]\s+|\s+-\s+", s, maxsplit=1)
            cur = {"name": m[0].strip(" :"), "description": m[1].strip() if len(m) > 1 else "", "technologies": [], "bullets": []}
            projs.append(cur)
        elif cur is not None:
            cur["description"] = (cur["description"] + " " + s).strip()
        else:
            cur = {"name": s[:60], "description": "", "technologies": [], "bullets": []}
            projs.append(cur)
    for p in projs:
        terms = find_terms(p["name"] + " " + p["description"] + " " + " ".join(p["bullets"]))
        p["technologies"] = list(dict.fromkeys(p["technologies"] + terms["programming_languages"] + terms["frameworks"] + terms["tools"]))
    return projs


def _parse_education(lines: list[str]) -> list[dict]:
    out: list[dict] = []
    cur: dict | None = None
    deg = re.compile(r"\b(b\.?\s?tech|m\.?\s?tech|b\.?\s?e\b|m\.?\s?e\b|b\.?\s?sc|m\.?\s?sc|bachelor|master|ph\.?d|doctorate|mba|bca|mca|diploma|associate|b\.?\s?a\b|m\.?\s?a\b|hsc|ssc|high school|secondary)", re.I)
    inst = re.compile(r"\b(university|college|institute|school|academy|iit|nit|polytechnic)\b", re.I)
    for raw in lines:
        s = _norm_line(raw)
        if not s:
            continue
        s = _BULLET.sub("", s)
        dates = ", ".join(re.findall(rf"{_DATE}(?:\s*(?:-|–|—|to)\s*(?:{_DATE}|present))?", s, re.I))
        has_deg, has_inst = bool(deg.search(s)), bool(inst.search(s))
        if has_deg or has_inst:
            if cur is None or (has_deg and cur["degree"]) or (has_inst and cur["institution"]):
                cur = {"degree": "", "institution": "", "dates": "", "details": ""}
                out.append(cur)
            parts = [p.strip() for p in re.split(r"\s*[|,–—]\s*|\s+-\s+|\s+at\s+", _RANGE.sub("", s)) if p.strip()]
            for p in parts:
                if inst.search(p) and not cur["institution"]:
                    cur["institution"] = p
                elif deg.search(p) and not cur["degree"]:
                    cur["degree"] = p
                elif p and not re.fullmatch(_DATE, p, re.I):
                    cur["details"] = (cur["details"] + "; " + p).strip("; ")
            cur["dates"] = cur["dates"] or dates
        elif cur is not None:
            if dates and not cur["dates"]:
                cur["dates"] = dates
            else:
                cur["details"] = (cur["details"] + "; " + s).strip("; ")
    return out


def _union_months(entries: list[dict]) -> float:
    spans = []
    for e in entries:
        a, b = _to_date(e["start"]), _to_date(e["end"], True)
        if a and b and b >= a:
            spans.append((a.year * 12 + a.month, b.year * 12 + b.month))
    spans.sort()
    total, cur_s, cur_e = 0, None, None
    for s, e in spans:
        if cur_e is None or s > cur_e:
            if cur_e is not None:
                total += cur_e - cur_s + 1
            cur_s, cur_e = s, e
        else:
            cur_e = max(cur_e, e)
    if cur_e is not None:
        total += cur_e - cur_s + 1
    return round(total / 12, 1)


def parse_resume(text: str) -> ParsedResume:
    text = text.replace("\r", "\n")
    lines = text.split("\n")
    head, secs, order = _split_sections(lines)
    r = ParsedResume(sections_found=order)
    r.name = _guess_name(head)
    r.summary = " ".join(_norm_line(_BULLET.sub("", l)) for l in secs.get("summary", []) if l.strip())[:1200]
    if not r.summary and not secs and len(head) > 3:
        r.warnings.append("No section headings recognised; using the whole text.")

    # skills ------------------------------------------------------------------------------
    skill_items: list[str] = []
    for ln in secs.get("skills", []):
        s = _norm_line(_BULLET.sub("", ln))
        if not s:
            continue
        s = re.sub(r"^[A-Za-z &/]{2,30}:\s*", "", s) if ":" in s[:32] else s
        skill_items += _split_list(s)
    all_text = "\n".join(secs.get("skills", [])) + "\n" + "\n".join(head)
    terms = find_terms(all_text)
    r.experience = _parse_experience(secs.get("experience", []))
    r.projects = _parse_projects(secs.get("projects", []))
    whole = find_terms(text)
    r.programming_languages = list(dict.fromkeys(terms["programming_languages"] + whole["programming_languages"]))
    r.frameworks = list(dict.fromkeys(terms["frameworks"] + whole["frameworks"]))
    r.tools = list(dict.fromkeys(terms["tools"] + whole["tools"]))
    r.technologies = list(dict.fromkeys(r.programming_languages + r.frameworks + r.tools))
    r.skills = list(dict.fromkeys(skill_items + r.technologies))[:80]
    if not secs.get("skills") and r.technologies:
        r.warnings.append("No Skills section; skills inferred from the text.")

    r.education = _parse_education(secs.get("education", []))
    certs = [_BULLET.sub("", _norm_line(l)) for l in secs.get("certifications", []) if l.strip()]
    r.certifications = [c for c in certs if len(c) > 2]
    ach = [_BULLET.sub("", _norm_line(l)) for l in secs.get("achievements", []) if l.strip()]
    r.achievements = ach + [a for e in r.experience for a in e["achievements"]]
    r.responsibilities = [b for e in r.experience for b in e["responsibilities"]]
    r.companies = list(dict.fromkeys(e["company"] for e in r.experience if e["company"]))
    r.job_titles = list(dict.fromkeys(e["title"] for e in r.experience if e["title"]))
    r.years_of_experience = _union_months(r.experience)
    if not r.years_of_experience:
        m = re.search(r"(\d{1,2})\s*\+?\s*(?:years?|yrs)", r.summary or text[:1500], re.I)
        if m:
            r.years_of_experience = float(m.group(1))
    if not r.experience:
        r.warnings.append("No work experience entries detected.")
    if not r.name:
        r.warnings.append("Could not determine the candidate's name.")
    if len(text.split()) < 30:
        r.warnings.append("Very little text in the resume.")
    return r
