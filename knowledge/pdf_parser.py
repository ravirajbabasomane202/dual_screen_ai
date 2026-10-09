"""Document parsing: PDF (pypdf), DOCX, TXT/MD -> pages of structural blocks."""
from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from app.errors import KnowledgeBaseError
from knowledge.structure import Block, clean_text, structure_text

SUPPORTED = (".pdf", ".txt", ".md", ".docx")
_PAGE_NO = re.compile(r"^\s*(?:page\s+)?\d{1,4}(?:\s*(?:/|of)\s*\d{1,4})?\s*$", re.I)


@dataclass
class ParsedPage:
    number: int
    blocks: list[Block]


@dataclass
class ParsedDocument:
    name: str
    path: str
    sha256: str
    size: int
    pages: list[ParsedPage] = field(default_factory=list)
    status: str = "ok"                 # ok | no_text
    warnings: list[str] = field(default_factory=list)


def file_sha256(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _edges(lines: list[str]) -> list[str]:
    nz = [ln for ln in lines if ln.strip()]
    return nz[:2] + nz[-2:]


def _strip_repeats(pages_text: list[str]) -> list[str]:
    """Drop bare page numbers and short lines repeated at the top/bottom of >=60% of pages (headers/footers)."""
    pages_lines = [t.split("\n") for t in pages_text]
    bad: set[str] = set()
    if len(pages_lines) >= 3:
        cnt: Counter = Counter()
        for lines in pages_lines:
            cnt.update({re.sub(r"\d+", "#", ln.strip()) for ln in _edges(lines) if 0 < len(ln.strip()) <= 60})
        bad = {k for k, v in cnt.items() if v >= 0.6 * len(pages_lines) and k.strip("# ")}
    out = []
    for lines in pages_lines:
        edge = set(_edges(lines))
        keep = [ln for ln in lines if not _PAGE_NO.match(ln)
                and not (ln in edge and re.sub(r"\d+", "#", ln.strip()) in bad)]
        out.append("\n".join(keep))
    return out


def _pdf_pages(path: Path, warnings: list[str]) -> list[str]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
    except Exception as e:
        raise KnowledgeBaseError(f"{path.name}: not a readable PDF ({type(e).__name__}: {e})") from e
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise KnowledgeBaseError(f"{path.name}: the PDF is password-protected")
    texts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            plain = page.extract_text() or ""
            try:                                   # layout mode only to recover column gaps for tables
                layout = page.extract_text(extraction_mode="layout") or ""
                if re.search(r"\S {3,}\S", layout) and len(layout.split()) >= len(plain.split()) * 0.8:
                    plain = "\n".join(ln.rstrip() for ln in layout.split("\n") if ln.strip())
            except Exception:
                pass
            texts.append(plain)
        except Exception as e:
            warnings.append(f"page {i} unreadable: {type(e).__name__}")
            texts.append("")
    return texts


def _docx_text(path: Path) -> list[str]:
    try:
        import docx
        d = docx.Document(str(path))
    except Exception as e:
        raise KnowledgeBaseError(f"{path.name}: not a readable DOCX ({e})") from e
    lines = []
    for p in d.paragraphs:
        style = (p.style.name or "") if p.style is not None else ""
        t = p.text.strip()
        if not t:
            lines.append("")
        elif style.lower().startswith("heading") or style == "Title":
            lines.append("# " + t)
        elif "list" in style.lower():
            lines.append("- " + t)
        else:
            lines.append(t)
    for tbl in d.tables:
        lines.append("")
        for row in tbl.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return ["\n".join(lines)]


def parse_document(path) -> ParsedDocument:
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in SUPPORTED:
        raise KnowledgeBaseError(f"{p.name}: unsupported file type")
    doc = ParsedDocument(p.name, str(p), file_sha256(p), p.stat().st_size)
    if ext == ".pdf":
        texts = _strip_repeats(_pdf_pages(p, doc.warnings))
    elif ext == ".docx":
        texts = _docx_text(p)
    else:
        raw = p.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texts = [raw.decode(enc)]
                break
            except UnicodeDecodeError:
                continue
    for i, t in enumerate(texts, 1):
        doc.pages.append(ParsedPage(i, structure_text(clean_text(t))))
    if not any(b.text.strip() for pg in doc.pages for b in pg.blocks):
        doc.status = "no_text"
        doc.warnings.append("no extractable text (scanned image PDF? OCR is not supported)")
    return doc
